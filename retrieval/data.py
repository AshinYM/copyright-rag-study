"""Validated local documents and immutable source offsets."""
from dataclasses import dataclass, asdict
import hashlib
import json
from pathlib import Path
import re
import unicodedata


def case_key(value):
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', value)).lower()


def case_numbers(value):
    return set(re.findall(r'\d{2,4}[가-힣]{1,4}\d+', case_key(value)))


@dataclass(frozen=True)
class Case:
    case_id: str
    title: str
    text: str
    decision_date: str = ''
    court: str = ''
    record_id: str = ''


@dataclass(frozen=True)
class Chunk:
    chunk_id: str
    case_id: str
    title: str
    start: int
    end: int
    text: str
    header: str
    record_id: str


def validate_cases(records):
    if not isinstance(records, list) or not records:
        raise ValueError('Cases must be a non-empty list')
    cases = []
    seen = set()
    for record in records:
        if not isinstance(record, dict) or set(record)-set(Case.__dataclass_fields__) or not {'case_id', 'title', 'text'} <= set(record):
            raise ValueError('Case schema: case_id, title, text, optional decision_date and court')
        if any(not isinstance(v, str) for v in record.values()):
            raise ValueError('Case values must be strings')
        if any(not record[k].strip() for k in ['case_id', 'title', 'text']):
            raise ValueError('Case ID, title and text must not be blank')
        key = case_key(record['case_id'])
        record_id = record.get('record_id') or key
        if record_id in seen:
            raise ValueError('Duplicate record identity; repeated case IDs require distinct record_id values')
        seen.add(record_id)
        # Preserve body exactly for source offsets; normalize only identity.
        cases.append(Case(**{**record, 'case_id': key, 'record_id': record_id}))
    return sorted(cases, key=lambda c: (c.case_id, c.record_id))


def load_cases(path):
    return validate_cases(json.loads(Path(path).read_text()))


def chunk_cases(cases, size=1000, overlap=150):
    if not isinstance(size, int) or not isinstance(overlap, int) or size <= 0 or not 0 <= overlap < size:
        raise ValueError('Require chunk_size > overlap >= 0')
    result = []
    for case in cases:
        header = f'{case.case_id} {case.title} {case.decision_date} {case.court}'
        for start in range(0, len(case.text), size-overlap):
            end = min(start+size, len(case.text))
            result.append(Chunk(f'{case.record_id}:{start}:{end}', case.case_id, case.title,
                                start, end, case.text[start:end], header, case.record_id))
            if end == len(case.text):
                break
    return result


def corpus_digest(cases):
    payload = json.dumps([asdict(c) for c in cases], ensure_ascii=False, sort_keys=True).encode()
    return hashlib.sha256(payload).hexdigest()


def save_new(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('x') as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write('\n')


def import_workbooks(corpus_path, questions_path, output):
    """Read Excel only. Preserve source body, ignore fully blank formatting rows."""
    from openpyxl import load_workbook
    output = Path(output)
    if output.exists():
        raise FileExistsError('Use a new import directory')
    required = ['사건명', '사건번호', '선고일자', '법원명', '판례내용']

    def read_rows(path, columns):
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = workbook.worksheets[0]
            rows = sheet.iter_rows(values_only=True)
            header = list(next(rows))
            if any(header.count(name) != 1 for name in columns):
                raise ValueError('Required workbook headers must appear exactly once')
            indices = [header.index(name) for name in columns]
            selected, empty = [], 0
            for row in rows:
                values = [row[i] if i < len(row) else None for i in indices]
                if not any(v is not None and str(v).strip() for v in values):
                    empty += 1
                    continue
                if any(v is None or not str(v).strip() for v in values):
                    raise ValueError('Incomplete non-empty workbook record')
                if any(isinstance(v, str) and v.startswith('=') for v in values):
                    raise ValueError('Formula cells require explicit review; never use stale cached values')
                selected.append([v.isoformat() if hasattr(v, 'isoformat') else str(v) for v in values])
            return selected, empty
        finally:
            workbook.close()

    rows, empty = read_rows(corpus_path, required)
    records, seen, removed_exact = [], set(), 0
    for row in rows:
        identity = hashlib.sha256(json.dumps(row, ensure_ascii=False).encode()).hexdigest()
        if identity in seen:
            removed_exact += 1
            continue
        seen.add(identity)
        records.append({**dict(zip(['title', 'case_id', 'decision_date', 'court', 'text'], row)),
                        'record_id': identity})
    cases = validate_cases(records)
    questions, question_empty = read_rows(questions_path, ['ID', 'Category', 'Question'])
    if len({r[0] for r in questions}) != len(questions):
        raise ValueError('Duplicate question IDs')
    pending = [{'id': row[0], 'category': row[1], 'question': row[2],
                'split': 'review', 'judgment': 'unjudged', 'relevant_cases': [],
                'evidence': [], 'reviewer': '', 'rationale': ''} for row in questions]
    from collections import Counter
    repeated = Counter(c.case_id for c in cases)
    summary = {'source_records': len(rows), 'preserved_records': len(cases),
               'unique_case_ids': len(repeated),
               'repeated_case_id_groups': sum(v > 1 for v in repeated.values()),
               'removed_exact_full_row_duplicates': removed_exact,
               'ignored_fully_empty_rows': empty,
               'questions': len(pending), 'ignored_empty_question_rows': question_empty,
               'official_qrels_available': False}
    save_new(output/'cases.json', [asdict(c) for c in cases])
    save_new(output/'questions-pending.json', {'kind': 'unjudged', 'items': pending})
    save_new(output/'import-audit.json', summary)
    return summary
